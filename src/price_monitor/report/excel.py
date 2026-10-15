"""Excel workbook with products, summary, changes, price history and mover charts."""

from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from price_monitor.models import ChangeKind
from price_monitor.report.data import ReportData

HEADER_FILL = PatternFill("solid", start_color="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SECTION_FONT = Font(bold=True, size=12, color="1F4E78")
TITLE_FONT = Font(bold=True, size=14, color="1F4E78")
LINK_FONT = Font(color="0563C1", underline="single")
PERCENT_FORMAT = "+0.0%;-0.0%;0.0%"
TIME_FORMAT = "yyyy-mm-dd hh:mm:ss"
_SYMBOLS = {"GBP": "£", "EUR": "€", "USD": "$", "JPY": "¥"}

KIND_LABELS = {
    ChangeKind.PRICE_CHANGE: "Price change",
    ChangeKind.OUT_OF_STOCK: "Out of stock",
    ChangeKind.BACK_IN_STOCK: "Back in stock",
    ChangeKind.NEW: "New",
    ChangeKind.REMOVED: "Removed",
}


def price_format(currency: str) -> str:
    symbol = _SYMBOLS.get(currency)
    return f'"{symbol}"#,##0.00' if symbol else f'#,##0.00 "{currency}"'


def _naive(value: datetime) -> datetime:
    """Excel has no time zones; values are written as naive UTC and labelled as such."""
    return value.replace(tzinfo=None, microsecond=0)


def _header(ws: Worksheet, labels: list[str]) -> None:
    """Append a styled header row."""
    ws.append(labels)
    row = ws.max_row
    for col in range(1, len(labels) + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _widths(ws: Worksheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _section(ws: Worksheet, title: str) -> None:
    ws.append([])
    ws.append([title])
    ws.cell(row=ws.max_row, column=1).font = SECTION_FONT


def _products_sheet(wb: Workbook, data: ReportData) -> None:
    ws = wb.active
    if ws is None:
        raise RuntimeError("workbook has no active sheet")
    ws.title = "Products"
    attribute_names = sorted({name for p in data.products for name in p.attributes})
    fixed = ["Site", "Key", "Title", "Category", "Price", "Currency", "Stock", "In stock", "URL"]
    _header(ws, fixed + [name.replace("_", " ").capitalize() for name in attribute_names])
    for product in data.products:
        in_stock = "" if product.in_stock is None else ("yes" if product.in_stock else "no")
        ws.append(
            [
                product.site,
                product.key,
                product.title,
                product.category,
                product.price,
                product.currency,
                product.stock,
                in_stock,
                product.url,
                *[product.attributes.get(name) for name in attribute_names],
            ]
        )
        row = ws.max_row
        ws.cell(row=row, column=5).number_format = price_format(product.currency)
        link = ws.cell(row=row, column=9)
        link.hyperlink = product.url
        link.font = LINK_FONT
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _widths(ws, [18, 20, 50, 20, 12, 10, 8, 10, 60] + [14] * len(attribute_names))


def _summary_sheet(wb: Workbook, data: ReportData) -> None:
    ws = wb.create_sheet("Summary")
    ws["A1"] = "Catalogue summary"
    ws["A1"].font = TITLE_FONT
    ws.append(["Generated (UTC)", _naive(data.generated_at)])
    ws.cell(row=ws.max_row, column=2).number_format = TIME_FORMAT
    ws.append(["Products tracked", len(data.products)])

    by_site: dict[str, list[Any]] = defaultdict(list)
    for p in data.products:
        by_site[p.site].append(p)
    _section(ws, "By site")
    _header(ws, ["Site", "Products", "Average price", "Currency", "In stock"])
    for site, items in sorted(by_site.items()):
        ws.append(
            [
                site,
                len(items),
                sum(i.price for i in items) / len(items),
                items[0].currency,
                sum(1 for i in items if i.in_stock),
            ]
        )
        ws.cell(row=ws.max_row, column=3).number_format = price_format(items[0].currency)

    _section(ws, "By category")
    _header(ws, ["Site", "Category", "Products", "Average price"])
    by_category: dict[tuple[str, str], list[Any]] = defaultdict(list)
    for p in data.products:
        by_category[(p.site, p.category or "(none)")].append(p)
    for (site, category), items in sorted(by_category.items()):
        ws.append([site, category, len(items), sum(i.price for i in items) / len(items)])
        ws.cell(row=ws.max_row, column=4).number_format = price_format(items[0].currency)

    _section(ws, "Runs")
    _header(ws, ["Run", "Site", "Started (UTC)", "Products", "Failures", "Simulated"])
    for run in data.runs:
        ws.append(
            [
                run.id,
                run.site,
                _naive(run.started_at),
                run.products_seen,
                run.failures,
                "yes" if run.simulated else "no",
            ]
        )
        ws.cell(row=ws.max_row, column=3).number_format = TIME_FORMAT
    _widths(ws, [22, 24, 20, 16, 12, 10])


def _changes_sheet(wb: Workbook, data: ReportData) -> None:
    ws = wb.create_sheet("Changes")
    ws["A1"] = "Changes in the latest run of each site"
    ws["A1"].font = TITLE_FONT
    if any(e.simulated for e in data.events):
        ws.append(
            ["Contains simulated changes (scrape --simulate-changes); not real site activity."]
        )
    if data.baseline_sites:
        ws.append(
            [f"Baseline only, nothing to compare against yet: {', '.join(data.baseline_sites)}."]
        )
    counts = Counter(e.kind for e in data.events)
    _section(ws, "Summary")
    for kind, label in KIND_LABELS.items():
        ws.append([label, counts.get(kind, 0)])
    _section(ws, f"Details ({len(data.events)})")
    if not data.events:
        ws.append(["No changes recorded."])
    else:
        _header(
            ws,
            ["Site", "Change", "Key", "Title", "Previous price", "New price", "Change %"],
        )
        for e in data.events:
            percent = (
                (e.new_price - e.old_price) / e.old_price
                if e.kind is ChangeKind.PRICE_CHANGE and e.old_price and e.new_price is not None
                else None
            )
            ws.append(
                [e.site, KIND_LABELS[e.kind], e.key, e.title, e.old_price, e.new_price, percent]
            )
            row = ws.max_row
            fmt = price_format(e.currency)
            ws.cell(row=row, column=5).number_format = fmt
            ws.cell(row=row, column=6).number_format = fmt
            ws.cell(row=row, column=7).number_format = PERCENT_FORMAT
    _widths(ws, [20, 16, 20, 50, 16, 14, 12])


def _history_sheet(wb: Workbook, data: ReportData) -> None:
    ws = wb.create_sheet("Price history")
    _header(
        ws, ["Site", "Key", "Title", "Observed (UTC)", "Price", "Currency", "Stock", "In stock"]
    )
    for product, point in data.history:
        in_stock = "" if point.in_stock is None else ("yes" if point.in_stock else "no")
        ws.append(
            [
                product.site,
                product.key,
                product.title,
                _naive(point.observed_at),
                point.price,
                product.currency,
                point.stock,
                in_stock,
            ]
        )
        ws.cell(row=ws.max_row, column=4).number_format = TIME_FORMAT
        ws.cell(row=ws.max_row, column=5).number_format = price_format(product.currency)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _widths(ws, [18, 20, 50, 18, 12, 10, 8, 10])


def _movers_sheet(wb: Workbook, data: ReportData) -> None:
    ws = wb.create_sheet("Top movers")
    ws["A1"] = "Largest price movers"
    ws["A1"].font = TITLE_FONT
    if not data.movers:
        ws.append(["No product has changed price yet; run the scraper again later."])
        return
    ws.append([])
    _header(ws, ["#", "Site", "Title", "First price", "Latest price", "Change %"])
    for rank, mover in enumerate(data.movers, start=1):
        fmt = price_format(mover.product.currency)
        ws.append(
            [
                rank,
                mover.product.site,
                mover.product.title,
                mover.first_price,
                mover.last_price,
                mover.percent / 100,
            ]
        )
        ws.cell(row=ws.max_row, column=4).number_format = fmt
        ws.cell(row=ws.max_row, column=5).number_format = fmt
        ws.cell(row=ws.max_row, column=6).number_format = PERCENT_FORMAT

    stamps = sorted({p.observed_at for m in data.movers for p in m.series})
    _section(ws, "Price series (chart data)")
    _header(ws, ["Observed (UTC)", *[f"#{i}" for i in range(1, len(data.movers) + 1)]])
    header_row = ws.max_row
    lookup = [{p.observed_at: p.price for p in m.series} for m in data.movers]
    for stamp in stamps:
        ws.append(
            [_naive(stamp).strftime("%Y-%m-%d %H:%M:%S"), *[series.get(stamp) for series in lookup]]
        )
    last_row = ws.max_row

    anchor_row = 3
    for index, mover in enumerate(data.movers):
        chart = LineChart()
        chart.title = mover.product.title[:60]
        chart.y_axis.title = mover.product.currency
        chart.x_axis.title = "Observed (UTC)"
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        chart.legend = None
        chart.height, chart.width = 7.5, 16
        values = Reference(ws, min_col=index + 2, min_row=header_row, max_row=last_row)
        chart.add_data(values, titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=header_row + 1, max_row=last_row))
        chart.series[0].marker.symbol = "circle"
        chart.series[0].smooth = False
        ws.add_chart(chart, f"I{anchor_row}")
        anchor_row += 16
    _widths(ws, [18, 18, 50, 12, 12, 12])


def write_excel(data: ReportData, path: Path) -> Path:
    wb = Workbook()
    _products_sheet(wb, data)
    _summary_sheet(wb, data)
    _changes_sheet(wb, data)
    _history_sheet(wb, data)
    _movers_sheet(wb, data)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
