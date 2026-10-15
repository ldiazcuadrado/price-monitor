from __future__ import annotations

import csv
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from price_monitor.models import Change, ChangeKind, Product
from price_monitor.report import load_report_data, write_excel, write_products_csv
from price_monitor.report.excel import price_format
from price_monitor.storage import Repository
from tests.helpers import charts

T0 = datetime(2026, 3, 1, 12, 0, tzinfo=UTC)


def product(key: str, price: float, site_currency: str = "GBP", **kw: Any) -> Product:
    return Product(key, f"Item {key}", price, site_currency, f"https://x/{key}", **kw)


def record_run(
    repo: Repository,
    hour: int,
    products: list[Product],
    changes: list[Change] | None = None,
    *,
    site: str = "shop",
    simulated: bool = False,
) -> int:
    when = T0 + timedelta(hours=hour)
    run_id = repo.start_run(site, when, page_limit=None, simulated=simulated)
    repo.record_observations(run_id, site, products, when)
    repo.add_events(run_id, site, changes or [], when)
    repo.finish_run(run_id, when, status="ok", products_seen=len(products))
    return run_id


def seed(repo: Repository) -> None:
    def catalogue(a: float, b: float, c: float) -> list[Product]:
        return [
            product("a", a, stock=3, in_stock=True, category="Poetry", attributes={"rating": 4}),
            product("b", b, category="Travel"),
            product("c", c, in_stock=False),
        ]

    record_run(repo, 0, catalogue(10.0, 20.0, 5.0))
    record_run(repo, 1, catalogue(12.0, 20.0, 5.0))
    changes = [
        Change(ChangeKind.PRICE_CHANGE, "a", "Item a", "https://x/a", "GBP", 12.0, 9.0),
        Change(ChangeKind.NEW, "c", "Item c", "https://x/c", "GBP", new_price=5.0),
    ]
    record_run(repo, 2, catalogue(9.0, 20.0, 5.0), changes, simulated=True)


def test_report_data_ranks_movers(repo: Repository) -> None:
    seed(repo)
    data = load_report_data(repo, top_n=5)
    assert [m.product.key for m in data.movers] == ["a"]
    assert data.movers[0].percent == -10.0
    assert len(data.history) == 9
    assert {e.kind for e in data.events} == {ChangeKind.PRICE_CHANGE, ChangeKind.NEW}
    assert data.baseline_sites == []


def test_excel_workbook_structure(repo: Repository, tmp_path: Path) -> None:
    seed(repo)
    path = write_excel(load_report_data(repo), tmp_path / "out" / "report.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == ["Products", "Summary", "Changes", "Price history", "Top movers"]

    products = wb["Products"]
    assert [c.value for c in products[1]][:9] == [
        "Site",
        "Key",
        "Title",
        "Category",
        "Price",
        "Currency",
        "Stock",
        "In stock",
        "URL",
    ]
    assert products.max_row == 4
    assert products["J1"].value == "Rating"
    first = next(r for r in products.iter_rows(min_row=2, values_only=True) if r[1] == "a")
    assert first[4] == 9.0 and first[9] == 4
    assert products["E2"].number_format == price_format("GBP")
    assert products["I2"].hyperlink is not None
    assert products.freeze_panes == "A2"

    history = wb["Price history"]
    assert history.max_row == 10
    assert [
        history.cell(row=r, column=5).value
        for r in range(2, 11)
        if history.cell(row=r, column=2).value == "a"
    ] == [10.0, 12.0, 9.0]


def test_excel_charts_for_top_movers(repo: Repository, tmp_path: Path) -> None:
    seed(repo)
    record_run(repo, 3, [product("b", 25.0), product("a", 9.0)], [])
    data = load_report_data(repo, top_n=2)
    wb = load_workbook(write_excel(data, tmp_path / "r.xlsx"))
    movers = wb["Top movers"]
    assert len(charts(movers)) == 2
    assert movers["C4"].value == "Item b" or movers["C5"].value == "Item b"
    assert all(chart.series for chart in charts(movers))


def test_excel_without_movers_has_no_charts(repo: Repository, tmp_path: Path) -> None:
    record_run(repo, 0, [product("a", 10.0)])
    wb = load_workbook(write_excel(load_report_data(repo), tmp_path / "r.xlsx"))
    assert charts(wb["Top movers"]) == []
    assert "No product has changed price" in wb["Top movers"]["A2"].value
    changes = [row[0].value for row in wb["Changes"].iter_rows() if row[0].value]
    assert any("Baseline only" in str(v) for v in changes)
    assert "No changes recorded." in changes


def test_changes_sheet_contents(repo: Repository, tmp_path: Path) -> None:
    seed(repo)
    wb = load_workbook(write_excel(load_report_data(repo), tmp_path / "r.xlsx"))
    rows = [[c.value for c in row] for row in wb["Changes"].iter_rows()]
    flat = [v for row in rows for v in row if v is not None]
    assert any("simulated" in str(v).lower() for v in flat)
    assert "Details (2)" in flat
    detail = next(r for r in rows if r[1] == "Price change")
    assert detail[:7] == ["shop", "Price change", "a", "Item a", 12.0, 9.0, -0.25]


def test_multi_currency_formats(repo: Repository, tmp_path: Path) -> None:
    record_run(repo, 0, [product("a", 10.0, "EUR")], site="eu")
    record_run(repo, 0, [product("a", 10.0, "SEK")], site="se")
    wb = load_workbook(write_excel(load_report_data(repo), tmp_path / "r.xlsx"))
    formats = {row[0].value: row[4].number_format for row in wb["Products"].iter_rows(min_row=2)}
    assert formats == {"eu": '"€"#,##0.00', "se": '#,##0.00 "SEK"'}


def test_site_filter(repo: Repository) -> None:
    record_run(repo, 0, [product("a", 1.0)], site="one")
    record_run(repo, 0, [product("b", 2.0)], site="two")
    data = load_report_data(repo, site="two")
    assert [p.key for p in data.products] == ["b"]
    assert [r.site for r in data.runs] == ["two"]


def test_csv_export(repo: Repository, tmp_path: Path) -> None:
    seed(repo)
    path = write_products_csv(load_report_data(repo), tmp_path / "p.csv")
    with path.open(encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    assert len(rows) == 3
    assert rows[0]["rating"] == "4"
    assert {r["key"] for r in rows} == {"a", "b", "c"}
    assert rows[1]["in_stock"] == ""


def test_summary_sheet_sections(repo: Repository, tmp_path: Path) -> None:
    seed(repo)
    wb = load_workbook(write_excel(load_report_data(repo), tmp_path / "r.xlsx"))
    cells = [row[0].value for row in wb["Summary"].iter_rows()]
    for title in ("Catalogue summary", "By site", "By category", "Runs"):
        assert title in cells
    header_row = cells.index("By site") + 1
    assert [c.value for c in wb["Summary"][header_row + 1]][:3] == [
        "Site",
        "Products",
        "Average price",
    ]
    site_row = [c.value for c in wb["Summary"][header_row + 2]]
    assert site_row[:2] == ["shop", 3]
    assert site_row[2] == pytest.approx((9 + 20 + 5) / 3)
